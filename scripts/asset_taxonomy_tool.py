#!/usr/bin/env python3
"""Asset taxonomy workbook generator and validation tool.

Usage:
  python scripts/asset_taxonomy_tool.py generate --output asset_taxonomy.xlsx
  python scripts/asset_taxonomy_tool.py validate \
      --taxonomy asset_taxonomy.xlsx \
      --raw site_register.xlsx \
      --output validation_report.csv
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable


TAXONOMY: dict[str, dict[str, list[str]]] = {
    "Building Services": {
        "HVAC Systems": ["Air Handling Unit", "Chiller", "Boiler", "Heat Pump"],
        "Electrical Distribution": ["Main Switchboard", "Distribution Board", "UPS"],
        "Water Services": ["Cold Water Booster Set", "Calorifier", "Water Heater"],
    },
    "Compliance Systems": {
        "Fire Safety": ["Fire Alarm Panel", "Smoke Detector", "Fire Extinguisher"],
        "Life Safety": ["Emergency Lighting", "Disabled Refuge Intercom"],
        "Security": ["CCTV Camera", "Access Control Panel", "Intruder Alarm"],
    },
    "Pool Plant": {
        "Filtration": ["Sand Filter", "Cartridge Filter", "Strainer"],
        "Disinfection": ["Dosing Pump", "UV Reactor", "Chemical Day Tank"],
        "Circulation": ["Pool Circulation Pump", "Balance Tank"],
    },
    "Operational Equipment": {
        "Leisure Equipment": ["Treadmill", "Exercise Bike", "Rowing Machine"],
        "Catering Equipment": ["Combi Oven", "Walk-In Freezer", "Dishwasher"],
        "Housekeeping Equipment": ["Ride-On Scrubber Dryer", "Commercial Washer"],
    },
    "Building Fabric": {
        "Roofs": ["Flat Roof Covering", "Pitched Roof Finish", "Rainwater Gutter"],
        "Internal Fabric": ["Fire Door", "Suspended Ceiling", "Partition Wall"],
        "Finishes": ["Floor Finish", "Wall Finish", "Decorative Coating"],
    },
    "External Infrastructure": {
        "Grounds": ["Perimeter Fence", "Car Park Surface", "Drainage Channel"],
        "External Lighting": ["Lighting Column", "Bollard Light"],
        "Utilities": ["Incoming Water Main", "Site Transformer", "Sewage Pump Station"],
    },
    "ICT & Systems": {
        "Network Infrastructure": ["Core Switch", "Wireless Access Point", "Router"],
        "Audio Visual": ["Digital Signage Player", "Public Address Amplifier"],
        "Building Systems": ["BMS Server", "BMS Controller", "Metering Gateway"],
    },
}

MANUFACTURERS = [
    "Daikin",
    "Mitsubishi Electric",
    "Grundfos",
    "Siemens",
    "Honeywell",
    "AstralPool",
    "Life Fitness",
]

CONTRACTORS = [
    "Apex Mechanical Services",
    "BlueWave Pool Engineering",
    "NorthStar Compliance Ltd",
    "Prime FM Solutions",
    "Urban Grid Electrical",
]

STATUS_VALUES = ["Active", "Inactive", "Commissioning", "Decommissioned"]


@dataclass
class ValidationIssue:
    row_number: int
    asset_id: str
    issue: str



def _style_header(ws, row: int = 1) -> None:
    from openpyxl.styles import Font, PatternFill

    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[row]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill


def _autosize(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[col_letter].width = min(max(12, max_len + 2), 50)


def create_workbook(output_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()

    ws_domains = wb.active
    ws_domains.title = "Domains"
    ws_domains.append(["Domain Code", "Domain Name", "Description"])

    ws_categories = wb.create_sheet("Categories")
    ws_categories.append(["Domain Name", "Category Code", "Category Name"])

    ws_types = wb.create_sheet("Types")
    ws_types.append(["Domain Name", "Category Name", "Type Code", "Type Name"])

    ws_manufacturers = wb.create_sheet("Manufacturers")
    ws_manufacturers.append(["Manufacturer Name"])

    ws_contractors = wb.create_sheet("Contractors")
    ws_contractors.append(["Contractor Name"])

    ws_upload = wb.create_sheet("Site Upload Template")
    ws_upload.append(
        [
            "Site",
            "Building",
            "Asset ID",
            "Asset Name",
            "Asset Domain",
            "Asset Category",
            "Asset Type",
            "Manufacturer",
            "Contractor",
            "Install Date",
            "Serial Number",
            "Status",
            "Notes",
        ]
    )

    ws_mapping = wb.create_sheet("Mapping Review")
    ws_mapping.append(
        [
            "Raw Domain",
            "Raw Category",
            "Raw Type",
            "Suggested Domain",
            "Suggested Category",
            "Suggested Type",
            "Compliance Status",
            "Reviewer Notes",
        ]
    )

    ws_lists = wb.create_sheet("Data Validation Lists")
    ws_lists.append(["Domains", "Categories", "Types", "Manufacturers", "Contractors", "Status"])

    domain_row = 2
    category_row = 2
    type_row = 2

    all_categories: list[str] = []
    all_types: list[str] = []

    for d_idx, (domain, categories) in enumerate(TAXONOMY.items(), start=1):
        ws_domains.append([f"DOM-{d_idx:02}", domain, f"{domain} asset domain"])
        ws_lists.cell(row=domain_row, column=1, value=domain)
        domain_row += 1

        for c_idx, (category, types) in enumerate(categories.items(), start=1):
            ws_categories.append([domain, f"CAT-{d_idx:02}-{c_idx:02}", category])
            all_categories.append(category)

            for t_idx, asset_type in enumerate(types, start=1):
                ws_types.append([domain, category, f"TYP-{d_idx:02}-{c_idx:02}-{t_idx:02}", asset_type])
                all_types.append(asset_type)

    for category in sorted(set(all_categories)):
        ws_lists.cell(row=category_row, column=2, value=category)
        category_row += 1

    for asset_type in sorted(set(all_types)):
        ws_lists.cell(row=type_row, column=3, value=asset_type)
        type_row += 1

    for idx, manufacturer in enumerate(MANUFACTURERS, start=2):
        ws_manufacturers.append([manufacturer])
        ws_lists.cell(row=idx, column=4, value=manufacturer)

    for idx, contractor in enumerate(CONTRACTORS, start=2):
        ws_contractors.append([contractor])
        ws_lists.cell(row=idx, column=5, value=contractor)

    for idx, status in enumerate(STATUS_VALUES, start=2):
        ws_lists.cell(row=idx, column=6, value=status)

    hierarchy_row = 1
    ws_lists.cell(row=hierarchy_row, column=8, value="Domain")
    ws_lists.cell(row=hierarchy_row, column=9, value="Category")
    ws_lists.cell(row=hierarchy_row, column=10, value="Type")

    for domain, categories in TAXONOMY.items():
        for category, types in categories.items():
            for asset_type in types:
                hierarchy_row += 1
                ws_lists.cell(row=hierarchy_row, column=8, value=domain)
                ws_lists.cell(row=hierarchy_row, column=9, value=category)
                ws_lists.cell(row=hierarchy_row, column=10, value=asset_type)

    wb.create_named_range("domain_list", ws_lists, f"$A$2:$A${domain_row - 1}")
    wb.create_named_range("category_list", ws_lists, f"$B$2:$B${category_row - 1}")
    wb.create_named_range("type_list", ws_lists, f"$C$2:$C${type_row - 1}")
    wb.create_named_range("manufacturer_list", ws_lists, f"$D$2:$D${1 + len(MANUFACTURERS)}")
    wb.create_named_range("contractor_list", ws_lists, f"$E$2:$E${1 + len(CONTRACTORS)}")
    wb.create_named_range("status_list", ws_lists, f"$F$2:$F${1 + len(STATUS_VALUES)}")

    max_rows = 2000
    domain_validation = DataValidation(type="list", formula1="=domain_list", allow_blank=True)
    category_validation = DataValidation(type="list", formula1="=category_list", allow_blank=True)
    type_validation = DataValidation(type="list", formula1="=type_list", allow_blank=True)
    manufacturer_validation = DataValidation(type="list", formula1="=manufacturer_list", allow_blank=True)
    contractor_validation = DataValidation(type="list", formula1="=contractor_list", allow_blank=True)
    status_validation = DataValidation(type="list", formula1="=status_list", allow_blank=True)

    ws_upload.add_data_validation(domain_validation)
    ws_upload.add_data_validation(category_validation)
    ws_upload.add_data_validation(type_validation)
    ws_upload.add_data_validation(manufacturer_validation)
    ws_upload.add_data_validation(contractor_validation)
    ws_upload.add_data_validation(status_validation)

    domain_validation.add(f"E2:E{max_rows}")
    category_validation.add(f"F2:F{max_rows}")
    type_validation.add(f"G2:G{max_rows}")
    manufacturer_validation.add(f"H2:H{max_rows}")
    contractor_validation.add(f"I2:I{max_rows}")
    status_validation.add(f"L2:L{max_rows}")

    for ws in [
        ws_domains,
        ws_categories,
        ws_types,
        ws_manufacturers,
        ws_contractors,
        ws_upload,
        ws_mapping,
        ws_lists,
    ]:
        _style_header(ws)
        _autosize(ws)
        ws.freeze_panes = "A2"

    ws_lists.sheet_state = "hidden"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def write_rules_document(path: Path) -> None:
    content = """# Asset Taxonomy Naming & Data Quality Rules

## 1) Hierarchy Standard
Use the fixed hierarchy for every asset record:

`Estate > Asset Domain > Asset Category > Asset Type > Individual Asset`

## 2) Asset ID Convention
- Pattern: `SITE-BLDG-DOM-CAT-####`
- Example: `LTR01-POOL-DOM03-CAT02-0007`
- Rules:
  - Uppercase alphanumeric with hyphen separators only.
  - Must be unique across the estate.
  - No spaces or special characters.

## 3) Name Formatting
- Domain / Category / Type names should use Title Case.
- Individual Asset Name should be descriptive and include duty/location where relevant.
  - Example: `Pool Circulation Pump 01 - Plantroom A`
- Avoid ambiguous abbreviations unless approved in an estate-wide abbreviation list.

## 4) Manufacturer and Contractor Standards
- Use the approved values maintained in workbook tabs:
  - `Manufacturers`
  - `Contractors`
- If not listed, add to review queue before using in production.

## 5) Mandatory Fields for Upload
Required in `Site Upload Template`:
- Site
- Asset ID
- Asset Name
- Asset Domain
- Asset Category
- Asset Type
- Status

## 6) Taxonomy Compliance Rules
A row is compliant only when:
1. Domain exists in approved taxonomy.
2. Category exists and belongs to the stated Domain.
3. Type exists and belongs to the stated Category + Domain.
4. Mandatory fields are populated.

## 7) Data Cleansing Rules
- Trim leading/trailing spaces from all text fields.
- Collapse multiple internal spaces to a single space.
- Standardise date fields to ISO format: `YYYY-MM-DD`.
- Replace placeholders like `TBC`, `N/A`, `UNKNOWN` with blank and route for review.
- Ensure statuses use approved controlled values only.

## 8) Mapping Review Process
Use the `Mapping Review` tab for any raw values that do not directly map.
- Record the raw value.
- Add the suggested mapped value.
- Mark status: `Compliant`, `Needs Review`, or `Rejected`.
- Capture reviewer notes and decision date.
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def _read_raw_rows(path: Path) -> Iterable[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            for row in reader:
                yield {k: (v or "").strip() for k, v in row.items() if k}
        return

    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]

    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {}
        for idx, header in enumerate(headers):
            row[header] = "" if idx >= len(values) or values[idx] is None else str(values[idx]).strip()
        yield row


def _load_taxonomy_sets(path: Path) -> tuple[set[str], set[tuple[str, str]], set[tuple[str, str, str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws_types = wb["Types"]

    domains: set[str] = set()
    domain_category: set[tuple[str, str]] = set()
    triplets: set[tuple[str, str, str]] = set()

    for domain, category, _, asset_type in ws_types.iter_rows(min_row=2, values_only=True):
        if not domain or not category or not asset_type:
            continue
        d = str(domain).strip()
        c = str(category).strip()
        t = str(asset_type).strip()
        domains.add(d)
        domain_category.add((d, c))
        triplets.add((d, c, t))

    return domains, domain_category, triplets


def validate_register(taxonomy_path: Path, raw_path: Path, output_path: Path) -> int:
    domains, domain_category, valid_triplets = _load_taxonomy_sets(taxonomy_path)

    required_fields = [
        "Site",
        "Asset ID",
        "Asset Name",
        "Asset Domain",
        "Asset Category",
        "Asset Type",
        "Status",
    ]

    issues: list[ValidationIssue] = []

    for row_num, row in enumerate(_read_raw_rows(raw_path), start=2):
        asset_id = row.get("Asset ID", "")

        missing = [f for f in required_fields if not row.get(f, "").strip()]
        if missing:
            issues.append(
                ValidationIssue(
                    row_number=row_num,
                    asset_id=asset_id,
                    issue=f"Missing required fields: {', '.join(missing)}",
                )
            )
            continue

        domain = row.get("Asset Domain", "").strip()
        category = row.get("Asset Category", "").strip()
        asset_type = row.get("Asset Type", "").strip()

        if domain not in domains:
            issues.append(ValidationIssue(row_num, asset_id, f"Unknown domain: {domain}"))
            continue

        if (domain, category) not in domain_category:
            issues.append(
                ValidationIssue(
                    row_num,
                    asset_id,
                    f"Category '{category}' is not valid for domain '{domain}'",
                )
            )
            continue

        if (domain, category, asset_type) not in valid_triplets:
            issues.append(
                ValidationIssue(
                    row_num,
                    asset_id,
                    (
                        f"Type '{asset_type}' is not valid for domain/category "
                        f"'{domain} / {category}'"
                    ),
                )
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["row_number", "asset_id", "issue"])
        for issue in issues:
            writer.writerow([issue.row_number, issue.asset_id, issue.issue])

    return len(issues)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Generate taxonomy workbook and validate raw registers.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    generate = subparsers.add_parser("generate", help="Generate a taxonomy workbook and rules document.")
    generate.add_argument("--output", default="output/asset_taxonomy_workbook.xlsx", type=Path)
    generate.add_argument("--rules", default="output/asset_naming_rules.md", type=Path)

    validate = subparsers.add_parser("validate", help="Validate a raw register against taxonomy workbook.")
    validate.add_argument("--taxonomy", required=True, type=Path)
    validate.add_argument("--raw", required=True, type=Path)
    validate.add_argument("--output", default="output/validation_report.csv", type=Path)

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.command == "generate":
        create_workbook(args.output)
        write_rules_document(args.rules)
        print(f"Workbook generated at {args.output}")
        print(f"Rules document generated at {args.rules}")
        return

    issue_count = validate_register(args.taxonomy, args.raw, args.output)
    print(f"Validation complete. Found {issue_count} non-compliant row(s). Report: {args.output}")


if __name__ == "__main__":
    main()
